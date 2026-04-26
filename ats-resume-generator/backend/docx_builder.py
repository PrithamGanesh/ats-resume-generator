"""
docx_builder.py
Builds an ATS-safe, single-column DOCX using python-docx.
No tables, no text boxes, no images — pure paragraph flow.
"""

import io
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Twips
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK  = RGBColor(0x1a, 0x1a, 0x2e)
GOLD  = RGBColor(0xc4, 0x9a, 0x1a)
GRAY  = RGBColor(0x44, 0x44, 0x44)
LGRAY = RGBColor(0x88, 0x88, 0x88)


def build_docx(data: dict, doc_type: str = "resume") -> io.BytesIO:
    doc = Document()

    # Page margins
    for sec in doc.sections:
        sec.top_margin    = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin   = Cm(2.0)
        sec.right_margin  = Cm(2.0)

    # Default font
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # Name
    _add_name(doc, data.get("name", "Your Name"))

    # Contact
    _add_contact(doc, data.get("contactLine", ""))

    # Summary
    if data.get("summary"):
        _section_heading(doc, "Professional Summary")
        _body(doc, data["summary"])

    # Core Competencies
    comps = data.get("coreCompetencies", [])
    if comps:
        _section_heading(doc, "Core Competencies")
        for i in range(0, len(comps), 3):
            _body(doc, "    •    ".join(comps[i:i+3]))

    # Experience
    for exp in data.get("experience", []):
        if not data.get("experience_heading_done"):
            _section_heading(doc, "Professional Experience")
            data["experience_heading_done"] = True
        _add_exp(doc, exp)

    # Education
    edu_list = data.get("education", [])
    if edu_list:
        _section_heading(doc, "Education")
        for ed in edu_list:
            _add_edu(doc, ed)

    # Certifications
    certs = [c for c in data.get("certifications", []) if c]
    if certs:
        _section_heading(doc, "Certifications")
        for c in certs:
            _bullet(doc, c)

    # Projects (CV only)
    projects = data.get("projects", [])
    if doc_type == "cv" and projects:
        _section_heading(doc, "Key Projects")
        for proj in projects:
            _bold_para(doc, proj.get("name", ""))
            for b in proj.get("bullets", []):
                _bullet(doc, b)

    # Additional Skills
    add_skills = data.get("additionalSkills", [])
    if add_skills:
        _section_heading(doc, "Additional Skills")
        _body(doc, "  •  ".join(add_skills))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ── Paragraph helpers ─────────────────────────────────────────────────────────

def _add_name(doc, name):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(name)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = DARK
    r.font.name = "Calibri"


def _add_contact(doc, contact):
    if not contact:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(contact)
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY
    r.font.name = "Calibri"
    _bottom_border(p, "C49A1A", "12")


def _section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(title.upper())
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = DARK
    r.font.name = "Calibri"
    _bottom_border(p, "C49A1A", "8")


def _body(doc, text, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.name = "Calibri"
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def _bold_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = "Calibri"
    r.font.color.rgb = DARK


def _bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Twips(180)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"▸  {text}")
    r.font.size = Pt(9.5)
    r.font.name = "Calibri"
    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def _add_exp(doc, exp):
    # Company + dates
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(1)
    r1 = p.add_run(exp.get("company", ""))
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = DARK
    r1.font.name = "Calibri"
    dates = exp.get("dates", "")
    if dates:
        r2 = p.add_run(f"    {dates}")
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = LGRAY
        r2.font.name = "Calibri"

    # Title + location
    title = exp.get("title", "")
    loc   = exp.get("location", "")
    if title or loc:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(3)
        line = title + (f"  —  {loc}" if loc else "")
        r = p2.add_run(line)
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = GRAY
        r.font.name = "Calibri"

    for b in exp.get("bullets", []):
        _bullet(doc, b)


def _add_edu(doc, ed):
    degree = ed.get("degree", "")
    field  = ed.get("field", "")
    label  = f"{degree}{' in ' + field if field else ''}"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(1)
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = DARK
    r.font.name = "Calibri"

    parts = [ed.get("institution",""), ed.get("year","")]
    if ed.get("gpa"):    parts.append(f"GPA: {ed['gpa']}")
    if ed.get("honors"): parts.append(ed["honors"])
    sub = "  |  ".join(p for p in parts if p)
    if sub:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(3)
        r2 = p2.add_run(sub)
        r2.font.size = Pt(9)
        r2.font.color.rgb = LGRAY
        r2.font.name = "Calibri"


def _bottom_border(para, color="C49A1A", size="8"):
    pPr  = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot  = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    size)
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), color)
    pBdr.append(bot)
    pPr.append(pBdr)
